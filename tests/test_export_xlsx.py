from openpyxl import load_workbook

from conftest import make_prompt, make_record, make_run_dir
from mizan.bank import Domain
from mizan.export.xlsx import write_workbook
from mizan.rundata import load_run


def _workbook(tmp_path, prompts, records):
    data = load_run(make_run_dir(tmp_path, prompts, records))
    out = write_workbook(data, tmp_path / "export.xlsx")
    return load_workbook(out)


def test_workbook_sheets_and_records(tmp_path):
    prompts = [make_prompt("a1", Domain.HISTORY, pair_id=None)]
    records = [make_record("a1", "openai/gpt-x", flag="answered")]
    wb = _workbook(tmp_path, prompts, records)

    assert wb.sheetnames == ["Records", "Model summary", "Domain breakdown", "Pairs"]

    ws = wb["Records"]
    header = [c.value for c in ws[1]]
    assert header[:6] == ["prompt_id", "domain", "pair_id", "query", "model", "flag"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2  # 1 prompt × 2 models, missing cell included
    captured = rows[0]
    assert captured[0] == "a1"
    assert captured[5] == "answered"
    assert captured[10] == "hello"
    assert rows[1][5] == "missing"
    assert rows[1][10] is None


def test_workbook_summary_sheets(tmp_path):
    prompts = [
        make_prompt("b1", Domain.COUNTERFACTUAL, pair_id="p1"),
        make_prompt("c1", Domain.COUNTERFACTUAL, pair_id="p1"),
    ]
    records = [
        make_record("b1", "openai/gpt-x", flag="answered"),
        make_record("c1", "openai/gpt-x", flag="refused"),
    ]
    wb = _workbook(tmp_path, prompts, records)

    summary = list(wb["Model summary"].iter_rows(min_row=2, values_only=True))
    gpt = summary[0]
    assert gpt[0] == "GPT-X"
    assert (gpt[1], gpt[3]) == (1, 1)  # answered, refused

    pairs = list(wb["Pairs"].iter_rows(min_row=2, values_only=True))
    assert len(pairs) == 2
    assert pairs[0][0] == "p1"
    assert (pairs[0][4], pairs[0][8]) == ("answered", "refused")  # a_flag, b_flag

    domains = list(wb["Domain breakdown"].iter_rows(min_row=2, values_only=True))
    assert domains[0][0] == "Counterfactual"


def test_workbook_truncates_long_response(tmp_path):
    prompts = [make_prompt("a1", Domain.HISTORY)]
    records = [make_record("a1", "openai/gpt-x", text="x" * 40_000)]
    wb = _workbook(tmp_path, prompts, records)

    text = list(wb["Records"].iter_rows(min_row=2, values_only=True))[0][10]
    assert len(text) < 32_767
    assert text.endswith("… [truncated]")
